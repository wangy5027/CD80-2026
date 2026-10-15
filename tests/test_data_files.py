"""Private data uploads and bounded ordinary-chat references."""

import io
import json
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from openpyxl import Workbook

from backend.app import create_app
from backend.app.modules.workbench.data_files import MAX_DATA_BYTES, build_data_info, data_context
from backend.app.modules.workbench.errors import GatewayError


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "持仓"
    sheet.append(["股票", "数量"])
    sheet.append(["甲", 12])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def workbook_with_rows(*rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class DataFilesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="cd80-personal-data-")
        self.config = {"TESTING": True, "DATA_DIR": self.temp.name, "GITEE_AI_API_KEY": "test-only-key",
                       "IMPORT_LEGACY_HISTORY": False, "CAPABILITY_SKIP_PUBLIC_SEED": True}
        self.app = create_app(self.config)
        self.service = self.app.extensions["workbench_service"]
        self.client = self.app.test_client()
        self.owner = self.register(self.client, "data-owner@example.invalid")

    def tearDown(self):
        self.service.close()
        self.temp.cleanup()

    @staticmethod
    def register(client, email):
        result = client.post("/desktop/api/register", json={"email": email, "password": "test-only-password"})
        assert result.status_code == 200, result.json
        return result.json["user"]["id"]

    def upload(self, name, data, *, purpose="data", client=None):
        return (client or self.client).post("/desktop/api/files", data={
            "file": (io.BytesIO(data), name), "purpose": purpose,
        })

    def test_formats_metadata_and_data_bytes_remain_private(self):
        cases = [
            ("sales.csv", "城市,金额\n成都,12\n".encode("utf-8"), "csv", "成都"),
            ("sales.tsv", "城市\t金额\n成都\t12\n".encode("gb18030"), "tsv", "成都"),
            ("sales.json", b'[{"city":"Chengdu","amount":12}]', "json", "Chengdu"),
            ("sales.xlsx", workbook_bytes(), "xlsx", "甲"),
        ]
        for name, data, expected_format, expected_value in cases:
            with self.subTest(name=name):
                result = self.upload(name, data)
                self.assertEqual(result.status_code, 201, result.json)
                metadata = result.json
                self.assertEqual(metadata["source"], "data_management")
                self.assertEqual(metadata["dataInfo"]["format"], expected_format)
                self.assertIn(expected_value, metadata["dataInfo"]["preview"])
                self.assertLessEqual(len(metadata["dataInfo"]["preview"]), 8000)
                if expected_format == "xlsx":
                    self.assertEqual(metadata["dataInfo"]["sheets"][0]["name"], "持仓")
                downloaded = self.client.get("/desktop/api/files/" + metadata["id"])
                self.assertEqual(downloaded.data, data)
                downloaded.close()

    def test_exact_size_boundary_and_over_limit(self):
        prefix = b'{"payload":"'
        suffix = b'"}'
        exact = prefix + b"x" * (MAX_DATA_BYTES - len(prefix) - len(suffix)) + suffix
        accepted = self.upload("large.json", exact)
        self.assertEqual(accepted.status_code, 201, accepted.json)
        self.assertEqual(accepted.json["size"], MAX_DATA_BYTES)
        self.assertTrue(accepted.json["dataInfo"]["truncated"])
        self.assertLessEqual(len(data_context(accepted.json, exact)), 12000)
        refused = self.upload("too-large.json", exact[:-1] + b"x" + exact[-1:])
        self.assertEqual(refused.status_code, 413)
        self.assertEqual(len(self.client.get("/desktop/api/state").json["files"]), 1)

    def test_empty_spoofed_and_unsafe_files_rejected(self):
        invalid = [
            ("empty.csv", b""),
            ("script.py", b"print('hello')"),
            ("fake.csv", b'{"hello":"json"}'),
            ("fake.xlsx", b"not a workbook"),
            ("corrupt.json", b'{"bad":'),
            ("nonfinite.json", b'{"amount": NaN}'),
            ("wrong.tsv", b"name,value\nA,1\n"),
            ("binary.csv", b"field\x00name\nfoo\n"),
        ]
        for name, data in invalid:
            with self.subTest(name=name):
                result = self.upload(name, data)
                self.assertEqual(result.status_code, 400, result.json)
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as packed:
            packed.writestr("[Content_Types].xml", "x")
            packed.writestr("xl/workbook.xml", "x")
            packed.writestr("../escape", "x")
        self.assertEqual(self.upload("unsafe.xlsx", archive.getvalue()).status_code, 400)
        self.assertEqual(self.client.get("/desktop/api/state").json["files"], [])

    def test_reupload_promotes_existing_file_and_persists_without_changing_bytes(self):
        data = b"name,value\nA,10\n"
        normal = self.upload("same.csv", data, purpose="")
        self.assertEqual(normal.status_code, 201, normal.json)
        self.assertNotIn("source", normal.json)
        promoted = self.upload("same.csv", data)
        self.assertEqual(promoted.status_code, 201, promoted.json)
        self.assertEqual(promoted.json["id"], normal.json["id"])
        self.assertEqual(promoted.json["source"], "data_management")
        self.assertEqual(self.upload("same.csv", data).json["id"], normal.json["id"])
        self.service.close()
        self.app = create_app(self.config)
        self.service = self.app.extensions["workbench_service"]
        restored = self.service.store.get(self.owner)["files"]
        self.assertEqual(len(restored), 1)
        self.assertEqual(restored[0]["dataInfo"]["columns"], ["name", "value"])
        self.assertEqual(self.client.get("/desktop/api/files/" + normal.json["id"]).data, data)

    def test_account_isolation_and_deleted_file_cannot_be_read(self):
        metadata = self.upload("private.csv", b"name,value\nA,10\n").json
        stranger = self.app.test_client()
        stranger_id = self.register(stranger, "stranger@example.invalid")
        self.assertNotEqual(stranger_id, self.owner)
        self.assertEqual(stranger.get("/desktop/api/state").json["files"], [])
        self.assertEqual(stranger.get("/desktop/api/files/" + metadata["id"]).status_code, 404)
        self.assertEqual(stranger.delete("/desktop/api/files/" + metadata["id"]).status_code, 404)
        unauthorized_chat = stranger.post("/desktop/api/chat", json={"message": "分析", "fileIds": [metadata["id"]]})
        self.assertEqual(unauthorized_chat.status_code, 404)
        self.assertEqual(self.client.delete("/desktop/api/files/" + metadata["id"]).status_code, 200)
        self.assertEqual(self.client.get("/desktop/api/files/" + metadata["id"]).status_code, 404)

    def test_current_and_historical_chat_include_bounded_csv_and_excel_content(self):
        csv_data = b"city,amount\nChengdu,12\n" + b"A,1\n" * 50000
        csv_meta = self.upload("large.csv", csv_data).json
        excel_meta = self.upload("holdings.xlsx", workbook_bytes()).json
        captured = {}

        def capture(context):
            captured.update(context)
            return iter(())

        auth = {"user": {"id": self.owner}}
        with patch.object(self.service, "_stream_chat", side_effect=capture):
            _stream, abandon = self.service.start_chat(auth, {"message": "请看资料", "fileIds": [csv_meta["id"], excel_meta["id"]]})
        abandon()
        self.assertIn("Chengdu", captured["query"])
        self.assertIn("甲", captured["query"])
        self.assertIn("有限预览", captured["query"])
        self.assertLess(len(captured["query"]), 26000)

        task = {"messages": [{"role": "user", "content": "请看资料", "attachments": captured["attachments"]}]}
        history = self.service._history_for_model(self.owner, task, self.service.store.get(self.owner))
        self.assertIn("Chengdu", history[0]["content"])
        self.assertIn("甲", history[0]["content"])
        self.assertLess(len(history[0]["content"]), 26000)

    def test_ordinary_upload_and_chat_limits_remain(self):
        ordinary = self.upload("notes.txt", b"a" * 180001, purpose="")
        self.assertEqual(ordinary.status_code, 201, ordinary.json)
        auth = {"user": {"id": self.owner}}
        with self.assertRaises(GatewayError) as raised:
            self.service.start_chat(auth, {"message": "阅读", "fileIds": [ordinary.json["id"]]})
        self.assertEqual(raised.exception.status, 400)

    def test_tampered_content_is_never_summarized(self):
        metadata = self.upload("check.csv", b"name,value\nA,10\n").json
        with self.assertRaises(GatewayError):
            data_context(metadata, b"name,value\nB,999\n")
        info = build_data_info("check.csv", b"name,value\nA,10\n")
        self.assertEqual(info["previewRows"], 1)

    def test_small_excel_uses_real_dimensions_without_padding(self):
        data = workbook_with_rows(("name", "amount"), ("A", 10), ("B", 20))
        result = self.upload("small.xlsx", data)
        self.assertEqual(result.status_code, 201, result.json)
        info = result.json["dataInfo"]
        self.assertEqual(info["columns"], ["name", "amount"])
        self.assertEqual(info["previewRows"], 2)
        self.assertFalse(info["truncated"])
        self.assertNotIn("null", info["preview"])
        self.assertEqual(info["preview"].count("\n"), 2)

    def test_excel_with_data_beyond_preview_window_marks_truncated(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "value"
        sheet["A100"] = 42
        output = io.BytesIO()
        workbook.save(output)
        result = self.upload("far-row.xlsx", output.getvalue())
        self.assertEqual(result.status_code, 201, result.json)
        info = result.json["dataInfo"]
        self.assertEqual(info["columns"], ["value"])
        self.assertEqual(info["previewRows"], 0)
        self.assertTrue(info["truncated"])

    def test_wide_table_is_accepted_and_preview_marks_truncation(self):
        header = [f"field_{index}" for index in range(70)]
        row = [str(index) for index in range(70)]
        csv_content = (",".join(header) + "\n" + ",".join(row) + "\n").encode()
        result = self.upload("wide.csv", csv_content)
        self.assertEqual(result.status_code, 201, result.json)
        info = result.json["dataInfo"]
        self.assertEqual(len(info["columns"]), 64)
        self.assertEqual(info["previewRows"], 1)
        self.assertTrue(info["truncated"])
        self.assertEqual(self.client.get("/desktop/api/files/" + result.json["id"]).data, csv_content)
        self.assertIn('"0"', data_context(result.json, csv_content, max_chars=4000))

    def test_large_field_and_long_names_are_valid_but_preview_is_truncated(self):
        long_field = b"x" * 200000
        field_result = self.upload("large-field.csv", b"value\n" + long_field + b"\n")
        self.assertEqual(field_result.status_code, 201, field_result.json)
        self.assertEqual(field_result.json["dataInfo"]["previewRows"], 1)
        self.assertTrue(field_result.json["dataInfo"]["truncated"])

        long_name = "heading" + "x" * 200
        name_result = self.upload("long-header.csv", (long_name + "\nA\n").encode())
        self.assertEqual(name_result.status_code, 201, name_result.json)
        self.assertTrue(name_result.json["dataInfo"]["columns"][0].endswith("…"))
        self.assertTrue(name_result.json["dataInfo"]["truncated"])

        excel_result = self.upload("long-cell.xlsx", workbook_with_rows(("name",), ("x" * 300,)))
        self.assertEqual(excel_result.status_code, 201, excel_result.json)
        self.assertTrue(excel_result.json["dataInfo"]["truncated"])

        wide_excel = self.upload("wide.xlsx", workbook_with_rows(tuple(f"c{i}" for i in range(70)), tuple(range(70))))
        self.assertEqual(wide_excel.status_code, 201, wide_excel.json)
        self.assertEqual(len(wide_excel.json["dataInfo"]["columns"]), 64)
        self.assertTrue(wide_excel.json["dataInfo"]["truncated"])


if __name__ == "__main__":
    unittest.main()
